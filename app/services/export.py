"""Render a Document's structured_data (+ WP-B extracted images) into
pdf/docx/xlsx bytes.

Best-effort STRUCTURED export. The docx/xlsx renderers here are the primary
output for those formats. `render_pdf` is now only a FALLBACK: WP-G made
PDF the visual-fidelity format, served by the /export router directly
(native-PDF pass-through, or the Tesseract searchable PDF) - `render_pdf`
runs only for pre-WP-G documents or when that artifact is missing.

Two things this module must get right, both consumed from data the pipeline
already produced - it changes no OCR/detection logic:

  * `structured_data["tables"]` entries have MORE THAN ONE SHAPE and must
    be branched on `source`:
      - "native_pdf"        -> {rows: list[list[str]], ...}  a real grid
      - "heuristic"         -> {rows: list[list[str]], ...}  a real grid
      - "ruled_line_region" -> {bbox, region_text: str}      a text blob,
        NO grid - render as labelled text, never as a fake table.
  * `images` (WP-B image_regions, with crop bytes resolved by the caller)
    are embedded. ALL of them - `region_type_guess` is only a caption hint,
    never a filter. WP-B's own principle: a silently dropped image is worse
    than a low-value one, and WP-B reported no "known-noise" guess class to
    exclude (its detector has ~0 false positives on text/tables; its weak
    spot is missed images and mislabelled types, not junk inclusions).
"""

import io
import logging
from typing import Any

from docx import Document as DocxDocument
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

logger = logging.getLogger("sense_tool.export")

_PDF_PAGE_WIDTH = 612  # US Letter, points
_PDF_PAGE_HEIGHT = 792
_PDF_FONT_SIZE = 10
_PDF_LINE_HEIGHT = 14
_PDF_LEFT_MARGIN = 40
_PDF_TOP_MARGIN = _PDF_PAGE_HEIGHT - 50
_PDF_BOTTOM_MARGIN = 40

_DOCX_MAX_IMG_WIDTH_IN = 6.0      # fits US-Letter with the default margins
_XLSX_MAX_IMG_WIDTH_PX = 600
_ASSUMED_IMG_DPI = 96.0


class ExportError(RuntimeError):
    """Raised when a document cannot be rendered in the requested format."""


# --------------------------------------------------------------- shared helpers
def _reencode_png(data: bytes) -> bytes:
    """Normalise arbitrary image bytes to a plain PNG via Pillow. Used as a
    fallback when a strict embedder (python-docx) rejects a technically-odd
    but decodable image."""
    from PIL import Image

    im = Image.open(io.BytesIO(data))
    im.load()
    buf = io.BytesIO()
    im.convert("RGB").save(buf, format="PNG")
    return buf.getvalue()


def _ordered_images(images: list[dict] | None) -> list[dict]:
    """Reading order: by page, then by the bbox's top y-coordinate. Exact
    pixel position is explicitly NOT the goal (WP-F) - this just keeps a
    page-1 letterhead ahead of a page-3 chart."""
    def key(m: dict):
        bbox = m.get("bbox") or []
        y = bbox[1] if len(bbox) >= 2 and isinstance(bbox[1], (int, float)) else 0
        return (m.get("page", 0) or 0, y)

    return sorted(images or [], key=key)


def _figure_caption(n: int, img: dict) -> str:
    guess = img.get("region_type_guess") or "unknown"
    src = img.get("source") or "?"
    page = (img.get("page", 0) or 0) + 1
    return f"Figure {n} (page {page}, {src}: {guess})"


# --------------------------------------------------------------- DOCX
def _docx_add_figures(doc, images: list[dict]) -> None:
    ordered = _ordered_images(images)
    if not ordered:
        return
    doc.add_heading("Figures", level=2)
    for n, img in enumerate(ordered, start=1):
        cap = doc.add_paragraph()
        cap.add_run(_figure_caption(n, img)).bold = True
        data = img.get("data")
        if not data:
            doc.add_paragraph(f"[Figure {n}: image data unavailable]")
            continue
        w_px = img.get("width") or 0
        kwargs: dict[str, Any] = {}
        if w_px:
            kwargs["width"] = Inches(min(_DOCX_MAX_IMG_WIDTH_IN, w_px / _ASSUMED_IMG_DPI))
        else:
            kwargs["width"] = Inches(_DOCX_MAX_IMG_WIDTH_IN)
        try:
            doc.add_picture(io.BytesIO(data), **kwargs)
        except Exception as first_exc:  # noqa: BLE001
            # python-docx's own image-header parser is stricter than Pillow
            # and rejects some valid embedded-PDF JPEGs (e.g. a signature
            # stamp in LabReport-1.pdf). Re-encode via Pillow and retry
            # before giving up - never silently drop the image (WP-B).
            try:
                data = _reencode_png(data)
                doc.add_picture(io.BytesIO(data), **kwargs)
            except Exception as exc:  # noqa: BLE001
                logger.warning("docx: could not embed figure %d: %s / %s", n, first_exc, exc)
                doc.add_paragraph(f"[Figure {n}: could not embed image ({exc or first_exc})]")


def _docx_add_tables(doc, tables: list[dict]) -> None:
    if not tables:
        return
    doc.add_heading("Tables", level=2)
    for idx, table in enumerate(tables, start=1):
        if table.get("source") == "ruled_line_region":
            # A detected ruled-table region that was OCR'd as one text block -
            # there is NO grid to render. Do not fabricate one.
            doc.add_heading(f"Table {idx} (detected, unstructured)", level=3)
            note = doc.add_paragraph()
            note.add_run(
                "Detected as a ruled table on a scanned page but not parsed "
                "into a grid; raw region text follows."
            ).italic = True
            doc.add_paragraph(str(table.get("region_text", "")).strip() or "(no text)")
            continue

        doc.add_heading(f"Table {idx}", level=3)
        rows = table.get("rows") or []
        if not rows:
            doc.add_paragraph("(empty table)")
            continue
        n_cols = max(len(r) for r in rows)
        docx_table = doc.add_table(rows=0, cols=n_cols)
        docx_table.style = "Table Grid"
        for row in rows:
            cells = docx_table.add_row().cells
            for i in range(n_cols):
                cells[i].text = str(row[i]) if i < len(row) else ""


def render_docx(source: str, structured_data: dict[str, Any] | None, images: list[dict] | None = None) -> bytes:
    structured_data = structured_data or {}
    images = images or []
    doc = DocxDocument()
    doc.add_heading(source, level=1)

    sections = structured_data.get("sections") or []
    if sections:
        doc.add_heading("Sections", level=2)
        for section in sections:
            doc.add_heading(str(section.get("title", "General")), level=3)
            content = str(section.get("content", ""))
            if content:
                doc.add_paragraph(content)

    _docx_add_figures(doc, images)

    tables = structured_data.get("tables") or []
    _docx_add_tables(doc, tables)

    if not sections and not tables and not images:
        doc.add_paragraph("(no structured data extracted)")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------- XLSX
def _xlsx_write_tables(ws, tables: list[dict]) -> None:
    for idx, table in enumerate(tables, start=1):
        if table.get("source") == "ruled_line_region":
            # XLSX is grid-shaped, but this table has no grid. Put the whole
            # region text in one wrapped, merged cell - do NOT fake columns.
            ws.append([f"Table {idx} - detected ruled region (unstructured text, not a grid)"])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            text_row = ws.max_row + 1
            ws.merge_cells(start_row=text_row, start_column=1, end_row=text_row, end_column=8)
            cell = ws.cell(row=text_row, column=1, value=str(table.get("region_text", "")).strip())
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.append([])
            ws.append([])
            continue

        ws.append([f"Table {idx} ({table.get('source', 'heuristic')})"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for row in table.get("rows") or []:
            ws.append([str(c) for c in row])
        ws.append([])


def _xlsx_add_figures_sheet(wb, images: list[dict], img_refs: list) -> None:
    """XLSX's grid model makes 'anchor an image next to its content'
    impractical, so - as WP-F allows - extracted images go on a dedicated
    "Figures" sheet, in reading order, each under a labelled header row."""
    ordered = _ordered_images(images)
    if not ordered:
        return
    try:
        from openpyxl.drawing.image import Image as XLImage
    except Exception as exc:  # noqa: BLE001 - Pillow missing etc.
        logger.warning("xlsx: image support unavailable (%s); listing figures as text", exc)
        XLImage = None

    ws = wb.create_sheet("Figures")
    ws.append(["Extracted images - one block per figure, in page/reading order."])
    ws.append([])
    row = ws.max_row + 1
    for n, img in enumerate(ordered, start=1):
        ws.cell(row=row, column=1, value=_figure_caption(n, img)).font = Font(bold=True)
        ws.cell(row=row + 1, column=1,
                value=f"size {img.get('width', '?')}x{img.get('height', '?')} "
                      f"format {img.get('format', '?')} bbox {img.get('bbox')}")
        data = img.get("data")
        anchor_row = row + 2
        placed = False
        if data and XLImage is not None:
            try:
                xi = XLImage(io.BytesIO(data))
                if xi.width and xi.width > _XLSX_MAX_IMG_WIDTH_PX:
                    scale = _XLSX_MAX_IMG_WIDTH_PX / xi.width
                    xi.width = int(xi.width * scale)
                    xi.height = int(xi.height * scale)
                ws.add_image(xi, f"A{anchor_row}")
                img_refs.append(xi)
                placed = True
            except Exception as exc:  # noqa: BLE001
                logger.warning("xlsx: could not embed figure %d: %s", n, exc)
        if not placed:
            ws.cell(row=anchor_row, column=1,
                    value="[image could not be embedded]" if data else "[image data unavailable]")
        # leave generous vertical space for the (unknown-height) picture
        row = anchor_row + 22


def render_xlsx(source: str, structured_data: dict[str, Any] | None, images: list[dict] | None = None) -> bytes:
    structured_data = structured_data or {}
    images = images or []
    wb = Workbook()
    _img_refs: list = []  # keep BytesIO-backed images alive until wb.save()

    sections_ws = wb.active
    sections_ws.title = "Sections"
    sections_ws.append(["Source", source])
    sections_ws.append([])
    sections_ws.append(["Title", "Content"])
    for section in structured_data.get("sections") or []:
        sections_ws.append([section.get("title", "General"), section.get("content", "")])

    tables = structured_data.get("tables") or []
    if tables:
        _xlsx_write_tables(wb.create_sheet("Tables"), tables)

    _xlsx_add_figures_sheet(wb, images, _img_refs)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------ PDF (structured-text FALLBACK)
# Used only when the /export router has no pixel-perfect PDF for a document
# (pre-WP-G rows, or a missing original/searchable artifact). The visual-
# fidelity PDF paths - native pass-through and the Tesseract searchable PDF -
# live in the router and the OCR worker step, not here.
def _structured_lines(
    source: str, structured_data: dict[str, Any] | None, images: list[dict] | None = None
) -> list[str]:
    """Flatten structured_data into text lines for the fallback plain-text
    PDF renderer - make sure the ruled-line region text and the image count
    are not silently lost when this fallback runs."""
    lines: list[str] = [f"Source: {source}", ""]
    structured_data = structured_data or {}
    images = images or []

    sections = structured_data.get("sections") or []
    if sections:
        lines.append("SECTIONS")
        for section in sections:
            lines.append(f"- {section.get('title', 'General')}")
            for content_line in str(section.get("content", "")).splitlines():
                if content_line.strip():
                    lines.append(f"    {content_line.strip()}")
        lines.append("")

    tables = structured_data.get("tables") or []
    if tables:
        lines.append("TABLES")
        for idx, table in enumerate(tables, start=1):
            if table.get("source") == "ruled_line_region":
                lines.append(f"- Table {idx} (detected, unstructured)")
                for t in str(table.get("region_text", "")).splitlines():
                    if t.strip():
                        lines.append(f"    {t.strip()}")
            else:
                lines.append(f"- Table {idx}")
                for row in table.get("rows", []):
                    lines.append("    " + " | ".join(str(c) for c in row))
        lines.append("")

    if images:
        lines.append("IMAGES")
        for n, img in enumerate(_ordered_images(images), start=1):
            lines.append(f"- {_figure_caption(n, img)}")
        lines.append("(image bytes are embedded in the docx / xlsx exports, not this PDF)")
        lines.append("")

    if not sections and not tables and not images:
        lines.append("(no structured data extracted)")

    return lines


def _pdf_escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def _wrap_line(line: str, max_chars: int = 100) -> list[str]:
    if len(line) <= max_chars:
        return [line]
    return [line[i : i + max_chars] for i in range(0, len(line), max_chars)]


def render_pdf(source: str, structured_data: dict[str, Any] | None, images: list[dict] | None = None) -> bytes:
    raw_lines = _structured_lines(source, structured_data, images)
    lines: list[str] = []
    for line in raw_lines:
        lines.extend(_wrap_line(line) or [""])
    if not lines:
        lines = [""]

    max_lines_per_page = int((_PDF_TOP_MARGIN - _PDF_BOTTOM_MARGIN) / _PDF_LINE_HEIGHT)

    writer = PdfWriter()

    for page_start in range(0, len(lines), max_lines_per_page):
        page_lines = lines[page_start : page_start + max_lines_per_page]
        page = writer.add_blank_page(width=_PDF_PAGE_WIDTH, height=_PDF_PAGE_HEIGHT)

        content_ops = [
            "BT",
            f"/F1 {_PDF_FONT_SIZE} Tf",
            f"{_PDF_LINE_HEIGHT} TL",
            f"{_PDF_LEFT_MARGIN} {_PDF_TOP_MARGIN} Td",
        ]
        for i, text_line in enumerate(page_lines):
            if i > 0:
                content_ops.append("T*")
            content_ops.append(f"({_pdf_escape(text_line)}) Tj")
        content_ops.append("ET")
        stream_data = "\n".join(content_ops).encode("latin-1", errors="replace")

        stream_obj = DecodedStreamObject()
        stream_obj.set_data(stream_data)
        stream_ref = writer._add_object(stream_obj)  # noqa: SLF001

        font_obj = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_ref = writer._add_object(font_obj)  # noqa: SLF001

        resources = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
        )
        page[NameObject("/Resources")] = resources
        page[NameObject("/Contents")] = stream_ref

    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


_RENDERERS = {
    "pdf": render_pdf,
    "docx": render_docx,
    "xlsx": render_xlsx,
}

_CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def render(
    format: str,
    source: str,
    structured_data: dict[str, Any] | None,
    images: list[dict] | None = None,
) -> bytes:
    """`images`: WP-B image_regions metadata dicts, each with the crop bytes
    resolved into an extra "data" key by the caller (the export router)."""
    try:
        renderer = _RENDERERS[format]
    except KeyError as exc:
        raise ExportError(f"Unsupported export format: {format!r}") from exc
    return renderer(source, structured_data, images)


def content_type_for(format: str) -> str:
    return _CONTENT_TYPES[format]
