import io
import uuid

import pytest
from docx import Document as DocxDocument
from openpyxl import load_workbook
from PIL import Image

from app.enums import DocumentStatus
from app.services import export
from app.storage import get_storage_backend

pytestmark = pytest.mark.asyncio(loop_scope="session")

STRUCTURED = {
    "sections": [{"title": "History", "content": "No prior admissions."}],
    "tables": [{"raw_lines": ["WBC | 7.2"], "rows": [["WBC", "7.2"]]}],
}


def _png(w: int = 24, h: int = 16, color=(180, 40, 40)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, format="PNG")
    return buf.getvalue()


async def test_export_pdf(client, make_document):
    document = await make_document(status=DocumentStatus.DONE, structured_data=STRUCTURED)
    response = await client.get(f"/documents/{document.id}/export", params={"format": "pdf"})
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


async def test_export_docx(client, make_document):
    document = await make_document(status=DocumentStatus.DONE, structured_data=STRUCTURED)
    response = await client.get(f"/documents/{document.id}/export", params={"format": "docx"})
    assert response.status_code == 200
    assert "wordprocessingml" in response.headers["content-type"]
    assert response.content.startswith(b"PK")


async def test_export_xlsx(client, make_document):
    document = await make_document(status=DocumentStatus.DONE, structured_data=STRUCTURED)
    response = await client.get(f"/documents/{document.id}/export", params={"format": "xlsx"})
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert response.content.startswith(b"PK")


async def test_export_rejects_invalid_format(client, make_document):
    document = await make_document(status=DocumentStatus.DONE, structured_data=STRUCTURED)
    response = await client.get(f"/documents/{document.id}/export", params={"format": "txt"})
    assert response.status_code == 422


async def test_export_404_when_missing(client):
    response = await client.get(f"/documents/{uuid.uuid4()}/export", params={"format": "pdf"})
    assert response.status_code == 404


async def test_export_409_when_not_done(client, make_document):
    document = await make_document(status=DocumentStatus.QUEUED)
    response = await client.get(f"/documents/{document.id}/export", params={"format": "pdf"})
    assert response.status_code == 409


async def test_export_409_when_failed(client, make_document):
    document = await make_document(status=DocumentStatus.FAILED, error_message="boom")
    response = await client.get(f"/documents/{document.id}/export", params={"format": "pdf"})
    assert response.status_code == 409


# --- WP-F: images flow from storage through the export endpoint ------------
# (pure-renderer table-shape / image tests live in test_export_render.py)


async def test_export_endpoint_embeds_stored_image_crops(client, make_document, db_session):
    """End-to-end: crop bytes live in the storage backend; the export
    endpoint resolves them and embeds them in the generated docx."""
    doc_id = uuid.uuid4()
    storage = get_storage_backend()
    key = f"{doc_id}/images/0.png"
    storage.save(key, _png(120, 90))

    from app.models import Document

    document = Document(
        id=doc_id, source="Clinic", status=DocumentStatus.DONE,
        raw_file_path=f"{doc_id}/raw.png", structured_data=STRUCTURED,
        image_regions=[{
            "storage_key": key, "bbox": [10, 20, 130, 110], "page": 0,
            "source": "embedded", "region_type_guess": "logo",
            "width": 120, "height": 90, "format": "png",
        }],
    )
    db_session.add(document)
    await db_session.commit()

    resp = await client.get(f"/documents/{doc_id}/export", params={"format": "docx"})
    assert resp.status_code == 200
    doc = DocxDocument(io.BytesIO(resp.content))
    assert len(doc.inline_shapes) == 1
    assert any("Figure 1" in p.text for p in doc.paragraphs)


async def test_export_endpoint_survives_missing_crop_file(client, make_document):
    """A dangling storage_key (crop file gone) must not 500 the export."""
    document = await make_document(
        status=DocumentStatus.DONE, structured_data=STRUCTURED,
        image_regions=[{"storage_key": "nope/missing.png", "bbox": [0, 0, 1, 1],
                        "page": 0, "source": "embedded", "region_type_guess": "logo"}],
    )
    resp = await client.get(f"/documents/{document.id}/export", params={"format": "docx"})
    assert resp.status_code == 200
    doc = DocxDocument(io.BytesIO(resp.content))
    assert any("image data unavailable" in p.text for p in doc.paragraphs)
