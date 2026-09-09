"""WP-F: pure export-renderer tests (no HTTP / DB) for the dual table shape
(native_pdf grid vs ruled_line_region text blob) and embedded images."""

import io

from docx import Document as DocxDocument
from openpyxl import load_workbook
from PIL import Image

from app.services import export


def _png(w: int = 24, h: int = 16, color=(180, 40, 40)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, format="PNG")
    return buf.getvalue()


MIXED_TABLES = {
    "sections": [{"title": "IMPRESSION", "content": "Mild anaemia."}],
    "tables": [
        {"source": "native_pdf", "raw_lines": ["Test | Result"],
         "rows": [["Test", "Result"], ["Haemoglobin", "11.4"]]},
        {"source": "ruled_line_region", "bbox": [30, 400, 700, 880],
         "region_text": "Type of Examination Result\nEYE 6/6\nB.P. 120/80"},
    ],
}

IMAGES = [
    {"storage_key": "d/images/0.jpg", "bbox": [25, 7, 585, 108], "page": 0,
     "source": "embedded", "region_type_guess": "logo", "width": 640, "height": 116,
     "format": "jpg", "data": _png(640, 116)},
    {"storage_key": "d/images/1.png", "bbox": [40, 620, 100, 680], "page": 1,
     "source": "detected", "region_type_guess": "chart", "width": 300, "height": 300,
     "format": "png", "data": _png(300, 300)},
]


def test_docx_native_grid_is_a_real_table_ruled_region_is_text():
    doc = DocxDocument(io.BytesIO(export.render_docx("Clinic", MIXED_TABLES, images=[])))

    assert len(doc.tables) == 1                            # only the native grid
    grid = [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert grid == [["Test", "Result"], ["Haemoglobin", "11.4"]]

    texts = "\n".join(p.text for p in doc.paragraphs)
    assert "Table 2 (detected, unstructured)" in texts
    assert "EYE 6/6" in texts and "B.P. 120/80" in texts


def test_docx_embeds_all_images_with_caption_hints():
    doc = DocxDocument(io.BytesIO(export.render_docx("Clinic", MIXED_TABLES, images=IMAGES)))
    assert len(doc.inline_shapes) == 2
    caps = "\n".join(p.text for p in doc.paragraphs)
    assert "Figure 1 (page 1, embedded: logo)" in caps     # ordered by (page, y)
    assert "Figure 2 (page 2, detected: chart)" in caps


def test_xlsx_native_grid_cells_and_ruled_region_merged_text():
    wb = load_workbook(io.BytesIO(export.render_xlsx("Clinic", MIXED_TABLES, images=[])))
    assert "Tables" in wb.sheetnames
    flat = [str(c) for r in wb["Tables"].iter_rows(values_only=True) for c in r if c]
    assert "Haemoglobin" in flat and "11.4" in flat
    assert any("detected ruled region" in c for c in flat)
    assert any("EYE 6/6" in c for c in flat)               # one blob, not fake columns


def test_xlsx_images_on_dedicated_figures_sheet():
    wb = load_workbook(io.BytesIO(export.render_xlsx("Clinic", MIXED_TABLES, images=IMAGES)))
    assert "Figures" in wb.sheetnames
    assert len(wb["Figures"]._images) == 2


def test_empty_document_exports_cleanly():
    empty = {"sections": [], "tables": []}

    dx = DocxDocument(io.BytesIO(export.render_docx("Clinic", empty, images=[])))
    assert len(dx.tables) == 0 and len(dx.inline_shapes) == 0
    assert any("no structured data" in p.text for p in dx.paragraphs)

    wb = load_workbook(io.BytesIO(export.render_xlsx("Clinic", empty, images=[])))
    assert wb.sheetnames == ["Sections"]                   # no empty Tables/Figures sheets

    assert export.render_pdf("Clinic", empty, images=[]).startswith(b"%PDF")
    assert export.render("docx", "Clinic", None, images=None).startswith(b"PK")
    assert export.render("xlsx", "Clinic", None).startswith(b"PK")


def test_docx_bad_image_bytes_do_not_break_export():
    imgs = [{"source": "detected", "region_type_guess": "photo", "page": 0,
             "bbox": [0, 0, 10, 10], "width": 100, "data": b"not-an-image"}]
    doc = DocxDocument(io.BytesIO(export.render_docx("Clinic", {"sections": [], "tables": []}, images=imgs)))
    assert len(doc.inline_shapes) == 0
    assert any("could not embed image" in p.text for p in doc.paragraphs)
